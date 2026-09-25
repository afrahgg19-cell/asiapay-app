from contextlib import contextmanager
from io import BytesIO
import os
import sqlite3
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

# إعدادات الصفحة
st.set_page_config(
    page_title="نظام إدارة المحفظة المالية الكبرى - ASIA PAY", layout="wide"
)

# --- دالة تطبيق تنسيق KPL (رمادي/رصاصي، حدود، خط 14) ---


def apply_kpl_styling_to_sheet(ws):
  """تطبيق تنسيق KPL رمادي/رصاصي مع حدود وخط 14 على الشيت"""
  header_fill = PatternFill(
      start_color='4A4A4A', end_color='4A4A4A', fill_type='solid'
  )
  header_font = Font(name='Calibri', size=14, bold=True, color='FFFFFF')

  row_fill_white = PatternFill(
      start_color='FFFFFF', end_color='FFFFFF', fill_type='solid'
  )
  row_fill_gray = PatternFill(
      start_color='F5F5F5', end_color='F5F5F5', fill_type='solid'
  )
  cell_font = Font(name='Calibri', size=14, color='000000')

  thin_side = Side(border_style='thin', color='D9D9D9')
  dark_side = Side(border_style='medium', color='595959')
  border_cell = Border(
      left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
  )
  border_header = Border(
      left=thin_side, right=thin_side, top=dark_side, bottom=dark_side
  )

  max_row = ws.max_row
  max_col = ws.max_column

  for row_idx in range(1, max_row + 1):
    is_header = row_idx == 1
    current_row_fill = header_fill if is_header else (
        row_fill_gray if row_idx % 2 == 0 else row_fill_white
    )

    for col_idx in range(1, max_col + 1):
      cell = ws.cell(row=row_idx, column=col_idx)
      cell.fill = current_row_fill
      cell.border = border_header if is_header else border_cell

      if is_header:
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal='center', vertical='center', wrap_text=True
        )
      else:
        cell.font = cell_font
        if isinstance(cell.value, (int, float)):
          cell.alignment = Alignment(horizontal='right', vertical='center')
          # تنسيق الأرقام تلقائياً بفواصل الآلاف (رقم حقيقي قابل للحساب)
          cell.number_format = '#,##0'
        else:
          cell.alignment = Alignment(horizontal='left', vertical='center')

  for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

  ws.freeze_panes = 'A2'


# --- دالة توحيد الشورت كود (تحل مشكلة تكرار نفس الكود بصيغ مختلفة) ---


def normalize_short_code(val):
  """
  توحيد شكل الشورت كود عشان نفس الكود ما يطلع مرتين بصيغ مختلفة
  (مثال: 1900 و 1900.0 و "1900 " تصير كلها "1900")
  """
  if val is None:
    return ""
  s = str(val).strip()
  if not s or s.lower() == "nan":
    return ""
  # شيل أي فواصل عشرية زايدة ناتجة عن قراءة الرقم كـ float (1900.0 -> 1900)
  if s.endswith(".0"):
    s = s[:-2]
  # شيل أي مسافات داخلية زايدة
  s = " ".join(s.split())
  return s.upper()


# --- دالة بحث مرنة عن عمود معين بغض النظر عن المسافات/حالة الأحرف ---


def find_col_flexible(df_columns, keywords):
  """
  تدور على أول عمود يحتوي أي كلمة من keywords، بعد تجاهل المسافات
  والشرطات السفلية وحالة الأحرف (كبيرة/صغيرة).
  مثال: "OrganizationArabicName" تُطابق الكلمة المفتاحية "arabicname"
  """
  for c in df_columns:
    c_norm = str(c).lower().replace(" ", "").replace("_", "").replace("-", "")
    for kw in keywords:
      if kw in c_norm:
        return c
  return None


# --- لوحة التحكم في الأعلى ---
st.markdown(
    "<h2 style='text-align: center; color: #1E3A8A;'>💰 نظام إدارة المحفظة"
    " المالية - ASIA PAY</h2>",
    unsafe_allow_html=True,
)

# استخدام الـ Tabs الأربعة العلوية
tab1, tab2, tab3, tab_kpi = st.tabs([
    "💳 محفظة ASIA PAY",
    "📊 المقارنة بين شهرين",
    "⭐ نسبة الإنجاز",
    "📈 KPI",
])

# ====================================================
# طبقة قاعدة البيانات الاحترافية (خاصة بتبويب المحفظة فقط)
# ====================================================
DB_FILE = "asia_pay_wallet.db"


@contextmanager
def get_db_connection():
  """
  Context manager احترافي للاتصال بقاعدة البيانات:
  - يفعّل WAL mode لتقليل تعارض القراءة/الكتابة وزيادة الموثوقية
  - يسوي commit تلقائي عند النجاح، و rollback تلقائي عند حدوث خطأ
  - يضمن إغلاق الاتصال دائماً حتى لو صار استثناء
  """
  conn = sqlite3.connect(DB_FILE, check_same_thread=False, timeout=10)
  try:
    conn.execute("PRAGMA journal_mode=WAL;")
    conn.execute("PRAGMA foreign_keys=ON;")
    conn.row_factory = sqlite3.Row
    yield conn
    conn.commit()
  except Exception:
    conn.rollback()
    raise
  finally:
    conn.close()


def init_db():
  """إنشاء الجدول والفهارس (Indexes) لتسريع الاستعلامات المتكررة"""
  try:
    with get_db_connection() as conn:
      conn.execute("""
          CREATE TABLE IF NOT EXISTS wallet_operations (
              id INTEGER PRIMARY KEY AUTOINCREMENT,
              timestamp TEXT NOT NULL,
              op_type TEXT NOT NULL,
              amount REAL NOT NULL CHECK (amount >= 0),
              details TEXT,
              payment_method TEXT,
              debt_status TEXT,
              remaining_balance REAL NOT NULL,
              created_at TEXT DEFAULT CURRENT_TIMESTAMP
          )
      """)
      conn.execute("""
          CREATE INDEX IF NOT EXISTS idx_wallet_timestamp
          ON wallet_operations (timestamp)
      """)
      conn.execute("""
          CREATE INDEX IF NOT EXISTS idx_wallet_op_type
          ON wallet_operations (op_type)
      """)
      conn.execute("""
          CREATE INDEX IF NOT EXISTS idx_wallet_debt_status
          ON wallet_operations (debt_status)
      """)
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء تهيئة قاعدة البيانات: {e}")


init_db()


def load_wallet_from_db():
  try:
    with get_db_connection() as conn:
      df = pd.read_sql(
          "SELECT * FROM wallet_operations ORDER BY id ASC", conn
      )
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء تحميل بيانات المحفظة: {e}")
    df = pd.DataFrame()

  if not df.empty:
    df = df.rename(
        columns={
            "timestamp": "التاريخ",
            "op_type": "نوع العملية",
            "amount": "المبلغ",
            "details": "التفاصيل / الجهة / السبب",
            "payment_method": "طريقة الدفع",
            "debt_status": "حالة الديون",
            "remaining_balance": "الباقي في المحفظة",
        }
    )
  else:
    df = pd.DataFrame(
        columns=[
            "id",
            "التاريخ",
            "نوع العملية",
            "المبلغ",
            "التفاصيل / الجهة / السبب",
            "طريقة الدفع",
            "حالة الديون",
            "الباقي في المحفظة",
        ]
    )
  return df


def get_latest_balance():
  try:
    with get_db_connection() as conn:
      row = conn.execute(
          "SELECT remaining_balance FROM wallet_operations"
          " ORDER BY id DESC LIMIT 1"
      ).fetchone()
    return row["remaining_balance"] if row else 0.0
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء قراءة الرصيد: {e}")
    return 0.0


def insert_operation(
    op_type, amount, details, payment_method, debt_status, new_balance
):
  """إدراج عملية جديدة بأمان (مع commit/rollback تلقائي عبر الـ context manager)"""
  try:
    with get_db_connection() as conn:
      conn.execute(
          """
              INSERT INTO wallet_operations
              (timestamp, op_type, amount, details, payment_method,
               debt_status, remaining_balance)
              VALUES (?, ?, ?, ?, ?, ?, ?)
          """,
          (
              str(pd.Timestamp.now()),
              op_type,
              amount,
              details,
              payment_method,
              debt_status,
              new_balance,
          ),
      )
    return True
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء حفظ العملية: {e}")
    return False


def update_operation(record_id, new_amount, new_details):
  try:
    with get_db_connection() as conn:
      conn.execute(
          "UPDATE wallet_operations SET amount = ?, details = ? WHERE id = ?",
          (new_amount, new_details, int(record_id)),
      )
    return True
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء تعديل السجل: {e}")
    return False


def delete_operation(record_id):
  try:
    with get_db_connection() as conn:
      conn.execute(
          "DELETE FROM wallet_operations WHERE id = ?", (int(record_id),)
      )
    return True
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء حذف السجل: {e}")
    return False


def mark_debt_paid(record_id):
  try:
    with get_db_connection() as conn:
      conn.execute(
          "UPDATE wallet_operations SET debt_status = 'تم التسديد'"
          " WHERE id = ?",
          (int(record_id),),
      )
    return True
  except Exception as e:
    st.error(f"⚠️ خطأ أثناء تحديث حالة التسديد: {e}")
    return False


# --- الحفاظ على حالة الجرد الكلي ومقارنة الشهور في الذاكرة ---
if "pivot_result" not in st.session_state:
  st.session_state["pivot_result"] = None
if "combined_df" not in st.session_state:
  st.session_state["combined_df"] = None
if "perf_summary" not in st.session_state:
  st.session_state["perf_summary"] = None


# ====================================================
# القسم الأول: محفظة ASIA PAY
# ====================================================
with tab1:
  st.markdown("### 💼 محفظة ASIA PAY (قاعدة بيانات دائمة)")
  st.markdown("---")

  df = load_wallet_from_db()
  last_balance = get_latest_balance()

  total_deposit = (
      df[df["نوع العملية"] == "إيداع للمحفظة"]["المبلغ"].sum()
      if not df.empty and "نوع العملية" in df.columns
      else 0.0
  )

  col1, col2, col3, col_target = st.columns(4)
  with col1:
    st.metric(
        label="الرصيد (الفعلي) الحالي في المحفظة", value=f"{last_balance:,.2f} د.ع"
    )
  with col2:
    st.metric(
        label="إجمالي مبالغ الإيداعات فقط", value=f"{total_deposit:,.2f} د.ع"
    )
  with col3:
    st.metric(
        label="إجمالي عدد الحركات المسجلة",
        value=str(len(df)) if not df.empty else "0",
    )
  with col_target:
    deposit_target_val = st.number_input(
        "🎯 تاركت الإيداع (Target)",
        value=st.session_state.get("deposit_target_val", 10000000.0),
        step=500000.0,
        format="%.2f",
        key="deposit_target_input",
    )
    st.session_state["deposit_target_val"] = deposit_target_val
    dep_progress = (
        (total_deposit / deposit_target_val) * 100.0
        if deposit_target_val > 0
        else 0.0
    )
    st.metric(
        label="نسبة إنجاز الإيداعات من التاركت", value=f"{dep_progress:,.2f}%"
    )

  st.markdown("---")

  c1, c2, c3 = st.columns(3)
  with c1:
    st.subheader("📥 إيداع للمحفظة")
    with st.form("deposit_form", clear_on_submit=True):
      deposit_amount = st.number_input(
          "المبلغ",
          value=None,
          min_value=0.0,
          step=1000.0,
          format="%.2f",
          key="dep_amt",
          placeholder="اكتب المبلغ هنا...",
      )
      deposit_reason = st.text_input("سبب الإيداع / اسم المودع", key="dep_res")
      submit_deposit = st.form_submit_button("حفظ الإيداع")
      if submit_deposit:
        amt_val = 0.0 if deposit_amount is None else float(deposit_amount)
        if amt_val > 0:
          current_bal = get_latest_balance()
          new_bal = current_bal + amt_val
          if insert_operation(
              "إيداع للمحفظة", amt_val, deposit_reason, "إيداع",
              "لا توجد", new_bal,
          ):
            st.success("تم حفظ الإيداع وتحديث الرصيد في قاعدة البيانات بنجاح!")
            st.rerun()
        else:
          st.warning("يرجى إدخال مبلغ صحيح أكبر من صفر.")

  with c2:
    st.subheader("📤 سحب كاش / مديونية")
    with st.form("withdraw_form", clear_on_submit=True):
      withdraw_amount = st.number_input(
          "المبلغ",
          value=None,
          min_value=0.0,
          step=1000.0,
          format="%.2f",
          key="wit_amt",
          placeholder="اكتب المبلغ هنا...",
      )
      withdraw_reason = st.text_input(
          "اسم المكاتب / السحب منه / المسؤول", key="wit_res"
      )
      payment_method = st.selectbox(
          "طريقة الدفع / الحالة", ["كاش", "ماستر كارد", "مديونية (دين)"]
      )
      submit_withdraw = st.form_submit_button("حفظ السحب")
      if submit_withdraw:
        amt_val = 0.0 if withdraw_amount is None else float(withdraw_amount)
        if amt_val > 0:
          current_bal = get_latest_balance()
          new_bal = current_bal - amt_val
          debt_status = (
              "غير مسدد (مديونية)"
              if payment_method == "مديونية (دين)"
              else "مكتمل"
          )
          if insert_operation(
              "سحب كاش", amt_val, withdraw_reason, payment_method,
              debt_status, new_bal,
          ):
            st.success("تم حفظ السحب وتحديث الرصيد في قاعدة البيانات بنجاح!")
            st.rerun()
        else:
          st.warning("يرجى إدخال مبلغ صحيح أكبر من صفر.")

  with c3:
    st.subheader("🔄 استرجاع مبالغ للمحفظة")
    with st.form("return_form", clear_on_submit=True):
      return_amount = st.number_input(
          "المبلغ الراجع",
          value=None,
          min_value=0.0,
          step=1000.0,
          format="%.2f",
          key="ret_amt",
          placeholder="اكتب المبلغ هنا...",
      )
      return_reason = st.text_input("سبب الاسترجاع / من الجهة", key="ret_res")
      submit_return = st.form_submit_button("إلغاء واسترجاع للمحفظة")
      if submit_return:
        amt_val = 0.0 if return_amount is None else float(return_amount)
        if amt_val > 0:
          current_bal = get_latest_balance()
          new_bal = current_bal + amt_val
          if insert_operation(
              "استرجاع للمحفظة", amt_val, return_reason, "استرجاع",
              "لا توجد", new_bal,
          ):
            st.success("تم استرجاع المبلغ وإضافته للمحفظة بنجاح!")
            st.rerun()
        else:
          st.warning("يرجى إدخال مبلغ صحيح أكبر من صفر.")

  st.markdown("---")
  st.subheader("📋 السجل التفصيلي للعمليات")
  if not df.empty:
    display_df = df.drop(columns=["id"], errors="ignore")
    st.dataframe(display_df, use_container_width=True)

    with st.expander("✏️ تعديل أو حذف عملية سابقة من السجل"):
      if "id" in df.columns:
        row_ids = df["id"].tolist()
        selected_id = st.selectbox(
            "اختر رقم السجل (ID) للتعديل أو الحذف:", row_ids
        )
        if selected_id:
          row_data = df[df["id"] == selected_id].iloc[0]
          with st.form("edit_row_form"):
            st.write(
                f"تعديل السجل ID: {selected_id} | التاريخ:"
                f" {row_data['التاريخ']}"
            )
            new_edit_amount = st.number_input(
                "تعديل المبلغ",
                value=float(row_data["المبلغ"]),
                step=1000.0,
                format="%.2f",
            )
            new_edit_reason = st.text_input(
                "تعديل التفاصيل / الجهة / السبب",
                value=str(row_data["التفاصيل / الجهة / السبب"]),
            )

            col_e1, col_e2 = st.columns(2)
            submit_edit = col_e1.form_submit_button("💾 حفظ التعديلات")
            submit_delete = col_e2.form_submit_button(
                "🗑️ حذف هذا السجل نهائياً"
            )

            if submit_edit:
              if update_operation(
                  selected_id, new_edit_amount, new_edit_reason
              ):
                st.success("تم تحديث السجل بنجاح!")
                st.rerun()

            if submit_delete:
              if delete_operation(selected_id):
                st.success("تم حذف السجل بنجاح!")
                st.rerun()
  else:
    st.info("لا توجد عمليات مسجلة حتى الآن.")

  st.markdown("---")
  st.subheader("📊 جرد الحسابات والإحصائيات الشاملة")
  if not df.empty:
    total_withdrawn = (
        df[df["نوع العملية"] == "سحب كاش"]["المبلغ"].sum()
        if "نوع العملية" in df.columns
        else 0.0
    )
    total_returned = (
        df[df["نوع العملية"] == "استرجاع للمحفظة"]["المبلغ"].sum()
        if "نوع العملية" in df.columns
        else 0.0
    )
    current_remaining = get_latest_balance()

    col_s1, col_s2, col_s3, col_s4 = st.columns(4)
    with col_s1:
      st.metric("إجمالي السحوبات", f"{total_withdrawn:,.2f} د.ع")
    with col_s2:
      st.metric("إجمالي الإيداعات", f"{total_deposit:,.2f} د.ع")
    with col_s3:
      st.metric("إجمالي المبالغ المسترجعة", f"{total_returned:,.2f} د.ع")
    with col_s4:
      st.metric("صافي رصيد المحفظة النهائي", f"{current_remaining:,.2f} د.ع")

  st.markdown("---")
  st.subheader("📋 قائمة الأشخاص والجهات المديونة (غير المسددة)")
  if "حالة الديون" in df.columns:
    debts_df = df[df["حالة الديون"] == "غير مسدد (مديونية)"]
    if not debts_df.empty:
      st.warning(f"تنبيه: لديك {len(debts_df)} مديونيات غير مسددة حالياً.")
      debt_list = []
      debt_map = {}
      for idx, row in debts_df.iterrows():
        label_text = f"ID ({row['id']}) - الجهة/الشخص: {row['التفاصيل / الجهة / السبب']} - المبلغ: {row['المبلغ']} د.ع"
        debt_list.append(label_text)
        debt_map[label_text] = row["id"]

      selected_debt_label = st.selectbox(
          "اختر المديونية لتسديدها:", debt_list
      )
      if st.button("✅ تم التسديد (تحديث وإزالة من المديونية)"):
        real_id = debt_map[selected_debt_label]
        if mark_debt_paid(real_id):
          st.success("تم تسديد المديونية وتحديث حالتها بنجاح!")
          st.rerun()
    else:
      st.info("ممتاز! لا توجد أي مديونيات معلقة حالياً، جميع الحسابات خالصة 🎉.")

# ====================================================
# القسم الثاني: المقارنة بين شهرين
# ====================================================
with tab2:
  st.markdown("### 📊 المقارنة بين أداء المكاتب بين شهرين")
  st.write("قم برفع ملف الشهر الأول والملف الثاني المقارن أدناه.")

  col_u1, col_u2 = st.columns(2)
  with col_u1:
    uploaded_file_8 = st.file_uploader(
        "اختر ملف الشهر الأول (Excel)", type=["xlsx", "xls"], key="file8"
    )
  with col_u2:
    uploaded_file_9 = st.file_uploader(
        "اختر ملف الشهر الثاني (Excel)", type=["xlsx", "xls"], key="file9"
    )

  if uploaded_file_8 is not None and uploaded_file_9 is not None:
    try:
      df8 = pd.read_excel(uploaded_file_8)
      df9 = pd.read_excel(uploaded_file_9)

      df8["Month"] = "الشهر الأول"
      df9["Month"] = "الشهر الثاني"

      combined_df = pd.concat([df8, df9], ignore_index=True)

      amt_candidates = [
          c
          for c in combined_df.columns
          if "amount" in str(c).lower() or "مبلغ" in str(c)
      ]
      amt_col = (
          amt_candidates[0] if amt_candidates else combined_df.columns[0]
      )

      def clean_amount(val):
        if pd.isna(val):
          return 0.0
        val_str = str(val).replace(",", "").strip()
        try:
          return float(val_str)
        except:
          return 0.0

      combined_df["Cleaned_Amount"] = combined_df[amt_col].apply(clean_amount)

      translation_dict = {
          (
              "Agency Commission Roll Up from Independent Store to Head Office"
          ): "ترحيل عمولات الوكالة من المتاجر المستقلة إلى الإدارة الرئيسية",
          "Auto Claw Back": "استرجاع تلقائي للأموال",
          "Commission Payment for Head Office": "دفع العمولات للإدارة الرئيسية",
          "Commission Payment for Independent Stores": (
              "دفع العمولات للمتاجر المستقلة"
          ),
          "Commission Roll Down for Independent Store": (
              "تنزيل العمولات للمتاجر المستقلة"
          ),
          "Customer Buy Goods Fee from Merchant": "أجور شراء بضائع من التاجر",
          "Customer Deposit at Agent": "إيداع نقدي للزبون لدى الوكيل",
          "Customer Withdraw at Agent": "سحب نقدي للزبون لدى الوكيل",
          "Organization Buy Airtime": "شراء رصيد / تعبئة من المؤسسة",
          "Organization Buy Electronic Vouchers": "شراء قسائم إلكترونية من المؤسسة",
          "Organization Deposit of Funds": "إيداع أموال للمؤسسة",
          (
              "Organization Inter Account Transfer - ORG to Agent"
          ): "تحويل بين حساب المؤسسة وحساب الوكيل",
          (
              "Organization Intra Account Transfer - Child to Child"
          ): "تحويل داخلي بين الفروع",
      }

      reason_col = (
          "Reason Type"
          if "Reason Type" in combined_df.columns
          else combined_df.columns[0]
      )
      combined_df["Arabic Translation"] = combined_df[reason_col].apply(
          lambda x: translation_dict.get(str(x), str(x))
      )

      code_col = (
          "Short Code"
          if "Short Code" in combined_df.columns
          else ("G" if "G" in combined_df.columns else combined_df.columns[0])
      )
      name_col = (
          "Arabic Name"
          if "Arabic Name" in combined_df.columns
          else (
              "F"
              if "F" in combined_df.columns
              else (
                  combined_df.columns
                  if len(combined_df.columns) > 1
                  else combined_df.columns[0]
              )
          )
      )

      combined_df["عدد حركات"] = 1

      pivot_result = combined_df.pivot_table(
          index=[code_col, name_col, reason_col, "Arabic Translation"],
          columns="Month",
          values=["Cleaned_Amount", "عدد حركات"],
          aggfunc={"Cleaned_Amount": "sum", "عدد حركات": "sum"},
          fill_value=0,
      ).reset_index()

      st.session_state["pivot_result"] = pivot_result
      st.session_state["combined_df"] = combined_df

      st.success("✅ تمت معالجة وحفظ المقارنة بين الشهرين بنجاح!")

    except Exception as e:
      st.error(f"⚠️ حدث خطأ أثناء المعالجة: {e}")

  if st.session_state["pivot_result"] is not None:
    st.subheader("📋 جدول مقارنة الجرد المحفوظ")
    st.dataframe(st.session_state["pivot_result"], use_container_width=True)

    output_filename = "Final_Inventory_Comparison_Report.xlsx"
    buffer_pivot = BytesIO()

    df_to_save_pivot = st.session_state["pivot_result"].copy()
    if isinstance(df_to_save_pivot.columns, pd.MultiIndex):
      df_to_save_pivot.columns = [
          "_".join([str(c) for c in col if col != ""])
          for col in df_to_save_pivot.columns
      ]

    with pd.ExcelWriter(buffer_pivot, engine="openpyxl") as writer:
      df_to_save_pivot.to_excel(writer, index=False, sheet_name="Comparison")

    # تطبيق تنسيق KPL على ملف البايفوت
    buffer_pivot.seek(0)
    wb_p = openpyxl.load_workbook(buffer_pivot)
    apply_kpl_styling_to_sheet(wb_p.active)
    buffer_pivot = BytesIO()
    wb_p.save(buffer_pivot)
    buffer_pivot.seek(0)

    st.download_button(
        label="📥 تحميل تقرير المقارنة (Excel)",
        data=buffer_pivot,
        file_name=output_filename,
        mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
  else:
    st.info("💡 يرجى رفع ملفات الشهرين في الأعلى لعرض وجرد البيانات.")

# ====================================================
# القسم الثالث: نسبة الإنجاز للمقارنة بين شهرين
# ====================================================
with tab3:
  st.markdown("### ⭐ نسبة الإنجاز والتقييم للمقارنة بين شهرين")
  st.write(
      "هذا القسم يعتمد على بيانات المقارنة بين الشهور لتقييم إنجاز المكاتب."
  )

  if (
      st.session_state["combined_df"] is not None
      and not st.session_state["combined_df"].empty
  ):
    df_combined = st.session_state["combined_df"]

    code_col = (
        "Short Code"
        if "Short Code" in df_combined.columns
        else ("G" if "G" in df_combined.columns else df_combined.columns[0])
    )
    name_col = (
        "Arabic Name"
        if "Arabic Name" in df_combined.columns
        else (
            "F"
            if "F" in df_combined.columns
            else (
                df_combined.columns
                if len(df_combined.columns) > 1
                else df_combined.columns[0]
            )
        )
    )

    if code_col in df_combined.columns and name_col in df_combined.columns:
      perf_summary = (
          df_combined.groupby([code_col, name_col])
          .agg(
              إجمالي_العمليات=("Cleaned_Amount", "count"),
              مجموع_المبالغ=("Cleaned_Amount", "sum"),
          )
          .reset_index()
      )

      target_benchmark = 10000000.0

      def calc_performance_and_progress(row):
        amt = row["مجموع_المبالغ"]
        progress_pct = min(100.0, (amt / target_benchmark) * 100.0)

        if amt > 5000000:
          perf_desc = "ممتاز (95%)"
          points = int(amt / 10000)
        elif amt > 2000000:
          perf_desc = "جيد جداً (85%)"
          points = int(amt / 10000)
        elif amt > 500050:
          perf_desc = "جيد (75%)"
          points = int(amt / 10000)
        else:
          perf_desc = "مقبول (60%)"
          points = int(amt / 10000)
        return pd.Series([perf_desc, progress_pct, points])

      perf_summary[[
          "نسبة الأداء",
          "نسبة الإنجاز (%)",
          "النقاط المكتسبة",
      ]] = perf_summary.apply(calc_performance_and_progress, axis=1)

      st.session_state["perf_summary"] = perf_summary
      st.success("✅ تم احتساب نسبة الإنجاز والتقييم للمكاتب!")
      st.dataframe(perf_summary, use_container_width=True)

      st.markdown("### 📈 مقارنة نسب الإنجاز للمكاتب")
      chart_df = perf_summary.set_index(name_col)["نسبة الإنجاز (%)"]
      st.bar_chart(chart_df)
    else:
      st.warning("⚠️ الأعمدة المطلوبة غير مطابقة.")
  else:
    st.info(
        "📌 يرجى رفع ملفات الشهرين في تبويب **(📊 المقارنة بين شهرين)** أولاً."
    )

# ====================================================
# التبويب الرابع: KPI (دمج الحركات + الإكسل الاختياري + ورقة Wallet report)
# ====================================================
with tab_kpi:
  st.markdown("### 📈 لوحة مؤشرات الأداء (KPI)")
  st.write(
      "1. رفـع ملف الإكسل الخاص بالحركات و Wallet report (إجباري).\n2. رفـع ملف المندوبين/الإكسل الاختياري لدمجه كلياً حسب الشورت كود."
  )

  col_k1, col_k2 = st.columns(2)
  with col_k1:
    kpi_uploaded_file = st.file_uploader(
        "اختر ملف الإكسل الخاص بالحركات و Wallet report (KPI)",
        type=["xlsx", "xls"],
        key="kpi_main_file_final_v7",
    )
  with col_k2:
    rep_uploaded_file = st.file_uploader(
        "اختر الإكسل الاختياري للمطوفة/المندوبين",
        type=["xlsx", "xls"],
        key="kpi_rep_file_final_v7",
    )

  if kpi_uploaded_file is not None:
    try:
      excel_obj = pd.ExcelFile(kpi_uploaded_file)
      sheet_names = excel_obj.sheet_names
      kpi_df = pd.read_excel(excel_obj, sheet_name=0)

      # --- استخراج رصيد المحفظة من Wallet report إن وجد ---
      wallet_balance_map = {}
      wallet_sheet_name = next(
          (s for s in sheet_names if "wallet" in s.lower()), None
      )
      if wallet_sheet_name:
        try:
          w_df = pd.read_excel(excel_obj, sheet_name=wallet_sheet_name)
          h_col_w = next(
              (c for c in w_df.columns if "accounttype" in str(c).lower()), None
          )
          r_col_w = next(
              (c for c in w_df.columns if "balance" in str(c).lower()), None
          )
          e_col_w = next(
              (
                  c
                  for c in w_df.columns
                  if str(c).strip().lower()
                  in ["shortcode", "short code", "short_code", "g", "e"]
              ),
              None,
          )
          if not e_col_w and len(w_df.columns) > 6:
            e_col_w = w_df.columns[6]
          if not h_col_w and len(w_df.columns) > 7:
            h_col_w = w_df.columns[7]
          if not r_col_w and len(w_df.columns) > 17:
            r_col_w = w_df.columns[17]

          if h_col_w and r_col_w and e_col_w:
            mask_h = (
                w_df[h_col_w].astype(str).str.strip()
                == "Organization E-Money Account"
            )
            filtered_w = w_df[mask_h].copy()

            def clean_balance_val(val):
              if pd.isna(val):
                return 0.0
              s = str(val).strip()
              if not s:
                return 0.0
              neg = False
              if s.startswith("(") and s.endswith(")"):
                neg = True
                s = s[1:-1].strip()
              s = s.replace(",", "")
              try:
                num = float(s)
                return -num if neg else num
              except ValueError:
                return 0.0

            filtered_w["cleaned_R"] = filtered_w[r_col_w].apply(
                clean_balance_val
            )
            filtered_w["key_clean"] = filtered_w[e_col_w].apply(
                normalize_short_code
            )
            wallet_balance_map = (
                filtered_w.groupby("key_clean")["cleaned_R"].sum().to_dict()
            )
        except Exception:
          pass

      def get_col_safe(preferred_name, fallback_idx, df_target):
        if preferred_name in df_target.columns:
          return preferred_name
        cols_local = [str(c).strip() for c in df_target.columns.tolist()]
        if len(cols_local) > fallback_idx:
          return df_target.columns[fallback_idx]
        return df_target.columns[0] if len(cols_local) > 0 else None

      g_col_name = get_col_safe("Short Code", 6, kpi_df)
      f_col_name = get_col_safe("Arabic Name", 5, kpi_df)
      b_col_name = get_col_safe("B", 1, kpi_df)
      c_col_name = get_col_safe("Reason Type", 2, kpi_df)
      t_col_name = get_col_safe("T", 19, kpi_df)

      work_kpi = pd.DataFrame()
      work_kpi["G_clean"] = (
          kpi_df[g_col_name].apply(normalize_short_code)
          if g_col_name in kpi_df.columns
          else pd.Series([""] * len(kpi_df))
      )
      # G_upper_key هو نفسه الكود الموحّد (نفس دالة التوحيد لكل المصادر)
      work_kpi["G_upper_key"] = work_kpi["G_clean"]
      work_kpi["F_clean"] = (
          kpi_df[f_col_name].astype(str).str.strip()
          if f_col_name in kpi_df.columns
          else pd.Series([""] * len(kpi_df))
      )
      work_kpi["B_clean"] = (
          kpi_df[b_col_name].astype(str).str.strip()
          if b_col_name in kpi_df.columns
          else pd.Series([""] * len(kpi_df))
      )
      # عمود Reason Type (C) - يُستخدم لتصنيف عمليات
      # "Organization Intra Account Transfer-Top to Child"
      work_kpi["C_clean"] = (
          kpi_df[c_col_name].astype(str).str.strip()
          if c_col_name in kpi_df.columns
          else pd.Series([""] * len(kpi_df))
      )

      raw_t_series = (
          kpi_df[t_col_name].astype(str)
          if t_col_name in kpi_df.columns
          else pd.Series(["0"] * len(kpi_df))
      )
      cleaned_t_numeric = (
          raw_t_series.str.replace(",", "", regex=False)
          .str.replace(" ", "", regex=False)
          .str.replace("$", "", regex=False)
      )
      work_kpi["T_num"] = pd.to_numeric(
          cleaned_t_numeric, errors="coerce"
      ).fillna(0.0)

      # شيل الصفوف اللي طلع كودها فاضي بعد التوحيد (صفوف فارغة/تالفة)
      work_kpi = work_kpi[work_kpi["G_upper_key"] != ""]

      target_ops = [
          "Merchant Payment",
          "Airtime Top-up",
          "Cash In",
          "Cash Out",
          "Bulk B2B Transfer",
          "Super Transaction",
          "E-money Deposit",
          "Electronic Vouchers",
      ]

      # --- تجميع مبالغ "Organization Intra Account Transfer-Top to Child"
      # مصنفة حسب الشورت كود الموحّد ---
      b2b_summary = (
          work_kpi[
              work_kpi["C_clean"].str.lower()
              == "organization intra account transfer-top to child"
          ]
          .groupby("G_upper_key")["T_num"]
          .sum()
          .to_dict()
      )

      kpi_grouped = {}
      for (g_v, f_v), grp in work_kpi.groupby(
          ["G_upper_key", "F_clean"], dropna=False
      ):
        g_str = str(g_v).strip()
        # لو نفس الكود عنده أكثر من اسم بالملف الأصلي، ندمج كل حركاته سوا
        if g_str in kpi_grouped:
          prev_f, prev_grp = kpi_grouped[g_str]
          kpi_grouped[g_str] = (
              prev_f or f_v,
              pd.concat([prev_grp, grp], ignore_index=True),
          )
        else:
          kpi_grouped[g_str] = (f_v, grp)

      all_short_codes = set(kpi_grouped.keys())

      # --- قراءة الدمج من الإكسل الاختياري ---
      opt_df = None
      opt_join_col = None
      opt_name_col = None
      opt_address_col = None
      opt_phone_col = None
      if rep_uploaded_file is not None:
        try:
          opt_df = pd.read_excel(rep_uploaded_file)

          # عمود الشورت كود بالملف الاختياري
          opt_join_col = find_col_flexible(
              opt_df.columns, ["shortcode", "كود"]
          )
          if opt_join_col is None and len(opt_df.columns) > 0:
            opt_join_col = opt_df.columns[0]

          # عمود الاسم بالملف الاختياري (يُستخدم كمصدر رئيسي للاسم) -
          # بحث مرن يطابق "OrganizationArabicName" أو "Arabic Name" أو غيرها
          opt_name_col = find_col_flexible(
              opt_df.columns,
              ["organizationarabicname", "arabicname", "الاسم", "name"],
          )
          # عمود العنوان بالملف الاختياري
          opt_address_col = find_col_flexible(
              opt_df.columns, ["address", "العنوان"]
          )
          # عمود رقم الهاتف بالملف الاختياري
          opt_phone_col = find_col_flexible(
              opt_df.columns, ["msisdn", "phone", "mobile", "هاتف"]
          )

          opt_df["_opt_key"] = opt_df[opt_join_col].apply(normalize_short_code)
          all_short_codes = all_short_codes | set(
              opt_df["_opt_key"].dropna().astype(str).tolist()
          )
        except Exception as e_opt:
          st.warning(f"⚠️ ملاحظة قراءة الإكسل الاختياري: {e_opt}")
          opt_df = None

      all_short_codes.discard("")

      opt_data_map = {}
      if opt_df is not None and "_opt_key" in opt_df.columns:
        other_cols = [c for c in opt_df.columns if c != "_opt_key"]
        for _, orow in opt_df.iterrows():
          o_k = str(orow["_opt_key"]).strip()
          if o_k and o_k != "NAN":
            opt_data_map[o_k] = {c: orow[c] for c in other_cols}

      kpi_rows_list = []
      for g_v in sorted(list(all_short_codes)):
        f_v = ""
        if g_v in kpi_grouped:
          f_val_found, grp = kpi_grouped[g_v]
          f_v = f_val_found
        else:
          grp = pd.DataFrame(
              columns=["G_clean", "F_clean", "B_clean", "C_clean", "T_num"]
          )

        # اسم المكتب: يُؤخذ أولاً من عمود الاسم بالإكسل الاختياري
        # (اللي تم لقطته بالبحث المرن)، ولو فاضي يرجع للاسم من الملف
        # الأصلي كحل احتياطي
        name_from_opt = ""
        if opt_name_col is not None:
          raw_name_val = opt_data_map.get(g_v, {}).get(opt_name_col, "")
          name_from_opt = (
              "" if pd.isna(raw_name_val) else str(raw_name_val).strip()
          )
        office_name = name_from_opt if name_from_opt else f_v

        # العنوان ورقم الهاتف: من الإكسل الاختياري عبر البحث المرن
        address_val = ""
        if opt_address_col is not None:
          raw_addr = opt_data_map.get(g_v, {}).get(opt_address_col, "")
          address_val = "" if pd.isna(raw_addr) else str(raw_addr).strip()

        phone_val = ""
        if opt_phone_col is not None:
          raw_phone = opt_data_map.get(g_v, {}).get(opt_phone_col, "")
          phone_val = "" if pd.isna(raw_phone) else str(raw_phone).strip()

        # مجموع "Organization Intra Account Transfer-Top to Child" لهذا
        # الشورت كود، مأخوذ من b2b_summary المُجهّز مسبقاً
        total_b2b_sum = b2b_summary.get(g_v, 0.0)

        if not grp.empty:
          high_t_count = int((grp["T_num"] > 4999).sum())
        else:
          high_t_count = 0

        # رقم حقيقي (float) بدل نص، حتى يبقى قابل للحساب والفلترة بالإكسل
        # مباشرة، وتنسيق الفواصل يصير من خصائص الإكسل (number_format)
        formatted_b2b = float(total_b2b_sum)

        w_bal = wallet_balance_map.get(g_v, 0.0)

        row_item = {
            "Short Code": g_v,
            "Organiztione Arabic name": office_name,
            "address": address_val,
            "msisdn": phone_val,
            "Busines to Business transfer": formatted_b2b,
            "حركه 100 الف": "Done" if total_b2b_sum > 99000 else "",
            "حركه 3 مليون": "Done" if total_b2b_sum > 2999000 else "",
            "اربع حركات": "Done" if high_t_count >= 4 else "",
            # رقم حقيقي هنا أيضاً لنفس السبب
            "رصيد المحفظة": (
                float(w_bal) if isinstance(w_bal, (int, float)) else w_bal
            ),
        }

        # دمج أي أعمدة إضافية أخرى من الإكسل الاختياري إن وجدت
        for c_k, c_v in opt_data_map.get(g_v, {}).items():
          if c_k not in row_item:
            row_item[c_k] = c_v

        for op in target_ops:
          if not grp.empty:
            count_val = grp["B_clean"].str.lower() == op.lower()
            row_item[f"عدد ({op})"] = int(count_val.sum())
          else:
            row_item[f"عدد ({op})"] = 0

        kpi_rows_list.append(row_item)

      final_kpi_table = pd.DataFrame(kpi_rows_list)

      # الترتيب النهائي المطلوب: كود المكتب - الاسم - العنوان - الهاتف -
      # مجموع المبالغ - الحركات - رصيد المحفظة - ثم بقية الأعمدة
      explicit_order = [
          "Short Code",
          "Organiztione Arabic name",
          "address",
          "msisdn",
          "Busines to Business transfer",
          "حركه 100 الف",
          "حركه 3 مليون",
          "اربع حركات",
          "رصيد المحفظة",
      ]
      existing_cols = [
          c for c in explicit_order if c in final_kpi_table.columns
      ]
      remaining_cols = [
          c for c in final_kpi_table.columns if c not in existing_cols
      ]
      final_kpi_table = final_kpi_table[existing_cols + remaining_cols]

      st.subheader(
          "📋 نتيجة تقرير الـ KPI (دمج شامل للحركات + الإكسل الاختياري +"
          " المحفظة بالترتيب المطلوب)"
      )
      st.dataframe(final_kpi_table, use_container_width=True)

      # --- تجهيز اسم ملف التحميل: اسم الملف الأصلي + تاريخ أمس (يوم-شهر) ---
      original_name_no_ext = os.path.splitext(kpi_uploaded_file.name)[0]
      yesterday_date = pd.Timestamp.now() - pd.Timedelta(days=1)
      date_suffix = f"{yesterday_date.day}-{yesterday_date.month}"
      out_kpi_name = f"{original_name_no_ext} {date_suffix}.xlsx"

      buffer_kpi = BytesIO()

      df_to_save_kpi = final_kpi_table.copy()
      if isinstance(df_to_save_kpi.columns, pd.MultiIndex):
        df_to_save_kpi.columns = [
            "_".join([str(c) for c in col if col != ""])
            for col in df_to_save_kpi.columns
        ]

      with pd.ExcelWriter(buffer_kpi, engine="openpyxl") as writer:
        df_to_save_kpi.to_excel(writer, index=False, sheet_name="KPI_Report")

      # تطبيق تنسيق KPL على تقرير الـ KPI النهائي
      buffer_kpi.seek(0)
      wb_k = openpyxl.load_workbook(buffer_kpi)
      apply_kpl_styling_to_sheet(wb_k.active)
      buffer_kpi = BytesIO()
      wb_k.save(buffer_kpi)
      buffer_kpi.seek(0)

      st.download_button(
          label="📥 تحميل تقرير KPI نهائي مدمج وشامل (Excel)",
          data=buffer_kpi,
          file_name=out_kpi_name,
          mime=(
              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
          ),
          key="download_kpi_excel_full_merged",
      )

    except Exception as err:
      st.error(f"⚠️ خطأ أثناء معالجة ملف الـ KPI: {err}")
  else:
    st.info("📌 يرجى رفع ملف الإكسل الرئيسي للـ KPI على الأقل لعرض النتائج.")
